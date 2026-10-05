"""Small files with hostile dimensions must have bounded processing work."""

import io
from unittest.mock import patch

from openpyxl import Workbook

from apps.documents.extraction import _extract_pdf, _extract_xlsx, extract_document_text


def test_spreadsheet_advertising_a_million_rows_is_bounded(settings):
    settings.XLSX_MAX_ROWS = 5
    settings.XLSX_MAX_COLUMNS = 3
    settings.XLSX_MAX_CELLS = 15
    workbook = Workbook()
    workbook.active["A1"] = "Visible opening text"
    workbook.active["XFD1048576"] = "Outside processing limit"
    payload = io.BytesIO()
    workbook.save(payload)
    result = _extract_xlsx(payload)
    assert "Visible opening text" in result.text
    assert "Outside processing limit" not in result.text
    assert any("processing limit" in note for note in result.notes)


def test_spreadsheet_total_cell_limit_applies_across_sheets(settings):
    settings.XLSX_MAX_ROWS = 10
    settings.XLSX_MAX_COLUMNS = 10
    settings.XLSX_MAX_CELLS = 2
    workbook = Workbook()
    workbook.active["A1"] = "First"
    workbook.create_sheet("Second")["A1"] = "Second text"
    workbook.create_sheet("Third")["A1"] = "Must not process"
    payload = io.BytesIO()
    workbook.save(payload)
    result = _extract_xlsx(payload)
    assert "First" in result.text and "Second text" in result.text
    assert "Must not process" not in result.text
    assert result.notes


def test_pdf_page_and_text_limits_preserve_original_page_count(settings):
    settings.OCR_MAX_PAGES = 2
    settings.OCR_MAX_CHARS = 60
    calls = []
    class Page:
        def extract_text(self):
            calls.append(1)
            return "Text from the document. " * 4
    class Reader:
        pages = [Page() for _ in range(5)]
    with patch("pypdf.PdfReader", return_value=Reader()):
        result = _extract_pdf(io.BytesIO())
    assert result.pages == 5
    assert len(result.text) <= 60
    assert len(calls) == 1
    assert any("first 2 pages" in note for note in result.notes)


def test_external_ocr_requires_an_explicit_opt_in(settings):
    settings.OCR_BACKEND = "ocr.space"
    with patch("apps.documents.extraction._ocr_space") as provider:
        result = extract_document_text(io.BytesIO(b"image"), "scan.png")
    provider.assert_not_called()
    assert result.status == "SKIPPED"


def test_unexpected_extraction_error_does_not_disclose_secret(settings):
    settings.OCR_SPACE_API_KEY = "test-secret-which-must-not-be-returned"
    with patch("apps.documents.extraction.extract_text_layer", side_effect=RuntimeError(settings.OCR_SPACE_API_KEY)):
        result = extract_document_text(io.BytesIO(b"test"), "note.txt")
    assert result.status == "FAILED"
    assert all(settings.OCR_SPACE_API_KEY not in note for note in result.notes)
