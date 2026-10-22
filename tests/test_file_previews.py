"""Private previews must not become a second, weaker download API."""

import io
from urllib.parse import parse_qs, urlparse

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from pypdf import PdfWriter

from apps.core.utils import sign_download
from apps.tracking.models import Attachment, RecordActivity
from apps.tracking.services import create_draft_record

pytestmark = pytest.mark.django_db


@pytest.fixture
def attachment(users, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=400)
    writer.write(output)
    record = create_draft_record(user=users['med'], subject='Preview source', instructions='Review')
    return Attachment.objects.create(
        record=record, original_name='signed.pdf', content_type='text/html',
        file=SimpleUploadedFile('signed.pdf', output.getvalue()), uploaded_by=users['med'],
    )


def token(kind, pk, user, purpose='preview'):
    return sign_download(f'file-version:{purpose}:{kind}:{pk}:{user.pk}')


def test_private_pdf_preview_keeps_status_and_original_bytes(client, users, attachment):
    client.force_login(users['med'])
    before_status = attachment.record.status
    before_activity = RecordActivity.objects.filter(record=attachment.record).count()
    response = client.get(f'/files/preview/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['med'])})
    assert response.status_code == 200
    assert response.context['preview_kind'] == 'pdf'
    assert response.context['version_number'] == 1
    assert 'Viewing Version 1 of 1' in response.content.decode()
    content_url = response.context['content_url']
    content = client.get(content_url)
    assert content.status_code == 200
    assert content['Content-Type'] == 'application/pdf'
    assert content['Cache-Control'] == 'private, no-store'
    assert content['X-Content-Type-Options'] == 'nosniff'
    with attachment.file.open('rb') as original:
        assert b''.join(content.streaming_content) == original.read()
    attachment.record.refresh_from_db()
    assert attachment.record.status == before_status
    assert RecordActivity.objects.filter(record=attachment.record).count() == before_activity


def test_preview_tokens_are_bound_to_kind_version_user_and_purpose(client, users, attachment):
    client.force_login(users['med'])
    base = f'/files/preview/tracking/{attachment.pk}/'
    assert client.get(base).status_code == 403
    assert client.get(base, {'token': token('tracking', attachment.pk, users['sup'])}).status_code == 403
    assert client.get(base, {'token': token('external', attachment.pk, users['med'])}).status_code == 403
    assert client.get(base, {'token': token('tracking', attachment.pk + 1, users['med'])}).status_code == 403
    preview = client.get(base, {'token': token('tracking', attachment.pk, users['med'])})
    assert preview.status_code == 200
    assert client.get(f'/files/download/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['med'])}).status_code == 403
    download = client.get(preview.context['download_url'])
    assert download.status_code == 200
    assert 'attachment;' in download['Content-Disposition']


def test_current_permissions_are_rechecked_on_each_content_request(client, users, attachment):
    client.force_login(users['sup'])
    response = client.get(f'/files/preview/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['sup'])})
    assert response.status_code == 404
    client.force_login(users['med'])
    preview = client.get(f'/files/preview/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['med'])})
    assert preview.status_code == 200
    users['med'].is_active = False
    users['med'].save(update_fields=['is_active'])
    denied = client.get(preview.context['content_url'])
    assert denied.status_code in {302, 403, 404}


def test_sandbox_is_the_only_embeddable_preview_document(client, users):
    client.force_login(users['med'])
    response = client.get('/files/preview-frame/')
    assert response.status_code == 200
    assert response['X-Frame-Options'] == 'SAMEORIGIN'
    policy = response['Content-Security-Policy']
    assert "script-src 'none'" in policy
    assert "connect-src 'none'" in policy
    assert "frame-ancestors 'self'" in policy
    assert b'<script' not in response.content


def test_content_url_is_not_a_storage_path_or_remote_url(client, users, attachment):
    client.force_login(users['med'])
    preview = client.get(f'/files/preview/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['med'])})
    assert preview.status_code == 200
    parsed = urlparse(preview.context['content_url'])
    assert not parsed.netloc
    assert parsed.path == f'/files/content/tracking/{attachment.pk}/'
    assert set(parse_qs(parsed.query)) == {'token'}


def test_tracking_latest_file_has_preview_and_version_history(client, users, attachment):
    client.force_login(users['med'])
    response = client.get(attachment.record.get_absolute_url())
    assert response.status_code == 200
    html = response.content.decode()
    assert f'/files/preview/tracking/{attachment.pk}/' in html
    assert 'Version 1' in html
    assert 'Upload new version' in html


def test_repository_append_endpoint_always_denies_even_empty_requests(client, users, settings, tmp_path):
    from apps.documents.models import Document

    settings.MEDIA_ROOT = tmp_path
    document = Document.objects.create(title='Frozen', office=users['med'].office, uploaded_by=users['med'])
    client.force_login(users['admin'])
    response = client.post(f'/documents/{document.pk}/files/')
    assert response.status_code == 403


def test_docx_preview_removes_external_and_active_parts():
    import zipfile

    from apps.core.file_previews import sanitized_docx

    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', '<document><altChunk/><object/><p>Safe</p></document>')
        archive.writestr('word/_rels/document.xml.rels', '<Relationships><Relationship TargetMode="External" Target="https://example.invalid/a.png" Type="image"/><Relationship Target="evil.html" Type="aFChunk"/></Relationships>')
        archive.writestr('word/evil.html', '<script>alert(1)</script>')
        archive.writestr('word/media/evil.svg', '<svg/>')
    with zipfile.ZipFile(io.BytesIO(sanitized_docx(output.getvalue()))) as archive:
        assert 'word/evil.html' not in archive.namelist()
        assert 'word/media/evil.svg' not in archive.namelist()
        assert b'https://' not in archive.read('word/_rels/document.xml.rels')
        assert b'altChunk' not in archive.read('word/document.xml')
        assert b'object' not in archive.read('word/document.xml')


def test_text_preview_escapes_html_and_discloses_truncation(client, users, attachment):
    item = Attachment.objects.create(record=attachment.record, uploaded_by=users['med'], original_name='notes.txt', file=SimpleUploadedFile('notes.txt', b'Notes <img src=x onerror=alert(1)>' + b'x' * 210000))
    client.force_login(users['med'])
    response = client.get(f'/files/preview/tracking/{item.pk}/', {'token': token('tracking', item.pk, users['med'])})
    assert response.status_code == 200
    assert b'&lt;img' in response.content
    assert b'Preview truncated' in response.content


def test_upload_version_post_conflicts_and_preserves_logical_count(client, users, attachment):
    client.force_login(users['med'])
    group = attachment.file_group_id
    url = f'/tracking/{attachment.record_id}/files/{group}/versions/'
    def submit(content):
        return client.post(url, {f'version-{group}-file': SimpleUploadedFile('updated.txt', content), f'version-{group}-expected_version': 1, f'version-{group}-change_summary': 'Signed outside DocTrack'})
    assert submit(b'Updated copy').status_code == 302
    assert submit(b'Stale copy').status_code == 409
    assert attachment.file_group.versions.count() == 2
    detail = client.get(attachment.record.get_absolute_url())
    assert len(detail.context['attachments']) == 1
    assert b'Version history (1 earlier)' in detail.content
    old_preview = client.get(f'/files/preview/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['med'])})
    assert b'Viewing Version 1 of 2' in old_preview.content


def test_repository_initial_multi_upload_then_metadata_and_history(client, users, settings, tmp_path):
    from apps.documents.models import Document

    settings.MEDIA_ROOT = tmp_path
    settings.ENABLE_BACKGROUND_TASKS = False
    client.force_login(users['med'])
    response = client.post('/documents/upload/', {'file': [SimpleUploadedFile('one.txt', b'Current one'), SimpleUploadedFile('two.txt', b'Current two')], 'office': users['med'].office_id, 'source': 'UPLOAD', 'ocr_language': 'auto'})
    assert response.status_code == 302
    document = Document.objects.get()
    assert document.files.count() == 2
    detail = client.get(document.get_absolute_url())
    assert detail.status_code == 200
    assert len(detail.context['file_rows']) == 2
    assert b'file contents' not in detail.content.lower() or b'read-only' in detail.content
    assert b'>Add files</button>' not in detail.content
    assert client.get(f'/documents/{document.pk}/edit/').status_code == 200
    version = document.files.first().versions.get()
    preview = client.get(f'/files/preview/repository/{version.pk}/', {'token': token('repository', version.pk, users['med'])})
    assert preview.status_code == 200


def test_malformed_docx_preview_is_readable_error_not_server_error(client, users, attachment):
    import zipfile
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w') as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', '<broken')
    item = Attachment.objects.create(record=attachment.record, original_name='bad.docx', file=SimpleUploadedFile('bad.docx', output.getvalue()), uploaded_by=users['med'])
    client.force_login(users['med'])
    preview = client.get(f'/files/preview/tracking/{item.pk}/', {'token': token('tracking', item.pk, users['med'])})
    response = client.get(preview.context['content_url'])
    assert response.status_code == 422
    assert b'original file is unchanged' in response.content


def test_empty_workbook_has_safe_preview_fallback(client, users, attachment):
    import re
    import zipfile

    from openpyxl import Workbook

    original = io.BytesIO()
    Workbook().save(original)
    malformed = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(malformed, 'w') as target:
        for name in source.namelist():
            data = source.read(name)
            if name == 'xl/workbook.xml':
                data = re.sub(b'<sheets>.*?</sheets>', b'<sheets/>', data)
            target.writestr(name, data)
    item = Attachment.objects.create(record=attachment.record, original_name='empty.xlsx', file=SimpleUploadedFile('empty.xlsx', malformed.getvalue()), uploaded_by=users['med'])
    client.force_login(users['med'])
    response = client.get(f'/files/preview/tracking/{item.pk}/', {'token': token('tracking', item.pk, users['med'])})
    assert response.status_code == 200
    assert b'Preview unavailable' in response.content


@pytest.mark.parametrize('broken_part', ['[Content_Types].xml', 'xl/workbook.xml'])
def test_malformed_workbook_xml_keeps_preview_fallback_and_original(client, users, attachment, broken_part):
    import zipfile

    from openpyxl import Workbook

    original = io.BytesIO()
    Workbook().save(original)
    malformed = io.BytesIO()
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(malformed, 'w') as target:
        for name in source.namelist():
            target.writestr(name, b'<broken' if name == broken_part else source.read(name))
    payload = malformed.getvalue()
    item = Attachment.objects.create(record=attachment.record, original_name='malformed.xlsx', file=SimpleUploadedFile('malformed.xlsx', payload), uploaded_by=users['med'])
    client.force_login(users['med'])
    before_status = attachment.record.status
    response = client.get(f'/files/preview/tracking/{item.pk}/', {'token': token('tracking', item.pk, users['med'])})
    assert response.status_code == 200
    assert b'Preview unavailable' in response.content
    download = client.get(response.context['download_url'])
    assert b''.join(download.streaming_content) == payload
    attachment.record.refresh_from_db()
    assert attachment.record.status == before_status


def test_workbook_chartsheets_are_not_treated_as_cell_worksheets():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    from apps.core.file_previews import tabular_preview

    book = Workbook()
    book.active.append(['Office', 'Count'])
    book.active.append(['MED', 1])
    chart = BarChart()
    chart.add_data(Reference(book.active, min_col=2, min_row=1, max_row=2))
    sheet = book.create_chartsheet('Chart')
    sheet.add_chart(chart)
    book._sheets.insert(0, book._sheets.pop())
    output = io.BytesIO()
    book.save(output)
    preview = tabular_preview(output.getvalue(), 'xlsx')
    assert preview['worksheets'] == ['Sheet']
    assert preview['rows'][1][0] == 'MED'


def test_required_initial_multi_upload_rejects_empty_list(client, users):
    client.force_login(users['med'])
    response = client.post('/documents/upload/', {'office': users['med'].office_id, 'source': 'UPLOAD', 'ocr_language': 'auto'})
    assert response.status_code == 200
    assert 'file' in response.context['form'].errors


def test_version_download_preserves_download_audit(client, users, attachment):
    from apps.core.models import AuditLog
    client.force_login(users['med'])
    preview = client.get(f'/files/preview/tracking/{attachment.pk}/', {'token': token('tracking', attachment.pk, users['med'])})
    response = client.get(preview.context['download_url'])
    assert response.status_code == 200
    assert AuditLog.objects.filter(actor=users['med'], action=AuditLog.Action.DOWNLOAD, target_id=str(attachment.record_id)).exists()


def test_image_preview_rejects_extreme_dimensions_before_browser_decode(client, users, attachment):
    import struct
    payload = b'\x89PNG\r\n\x1a\n' + b'\x00\x00\x00\rIHDR' + struct.pack('>II', 100000, 100000) + b'\x08\x02\x00\x00\x00'
    item = Attachment.objects.create(record=attachment.record, original_name='huge.png', file=SimpleUploadedFile('huge.png', payload), uploaded_by=users['med'])
    client.force_login(users['med'])
    preview = client.get(f'/files/preview/tracking/{item.pk}/', {'token': token('tracking', item.pk, users['med'])})
    response = client.get(preview.context['content_url'])
    assert response.status_code == 422


def test_version_upload_forms_keep_internal_revision_fields_hidden(client, users, attachment):
    client.force_login(users['med'])
    response = client.get(attachment.record.get_absolute_url())
    assert b'>Expected version' not in response.content
    assert b'>Expected revision' not in response.content


def test_mislabeled_embedded_raster_cannot_bypass_pixel_limit():
    from django.core.exceptions import ValidationError

    from apps.core.file_previews import image_pixels
    # Its APP bytes resemble a 1x1 GIF header, but the true JPEG SOF is huge.
    payload = b'\xff\xd8\xff\xe0\x00\x06\x01\x00\x01\x00\xff\xc0\x00\x0b\x08\xea\x60\xea\x60\x01\x01\x11\x00'
    with pytest.raises(ValidationError):
        image_pixels(payload, 'gif')
