"""Native readers preserve source order, late content and coverage uncertainty."""

import io
from dataclasses import replace
from unittest.mock import patch

import pytest
from docx import Document as WordDocument
from openpyxl import Workbook
from PIL import Image

from apps.documents.processing_contracts import ExtractionPolicy
from tests.intelligence_fixtures import image_bytes
from tests.test_pdf_coverage import reader


def test_word_body_order_header_footer_and_embedded_image():
    word = WordDocument()
    word.sections[0].header.paragraphs[0].text = "Header policy"
    word.add_paragraph("Opening paragraph")
    word.add_table(rows=1, cols=1).cell(0, 0).text = "Middle table"
    word.add_paragraph("Late appendix phrase")
    word.add_picture(io.BytesIO(image_bytes()))
    word.sections[0].footer.paragraphs[0].text = "Footer reference"
    stream = io.BytesIO()
    word.save(stream)
    batch = reader()(stream, "ordered.docx", policy=ExtractionPolicy())
    text = "\n".join(unit.normalized_text for unit in batch.units)
    assert text.index("Opening paragraph") < text.index("Middle table") < text.index("Late appendix phrase")
    assert "Header policy" in text and "Footer reference" in text
    assert any(unit.kind == "embedded_image" for unit in batch.units)
    assert batch.coverage != "complete"
    assert not any("page:" in unit.locator_key for unit in batch.units)


def test_many_small_word_blocks_stop_at_an_explicit_unit_budget():
    word = WordDocument()
    for i in range(2100):
        word.add_paragraph(f"Authored paragraph {i}")
    output = io.BytesIO()
    word.save(output)
    batch = reader()(output, "many-blocks.docx", policy=ExtractionPolicy())
    assert len(batch.units) <= 2000
    assert batch.expected_units is None
    assert batch.coverage != "complete" and batch.omitted_reasons
    assert "Authored paragraph 0" in batch.units[0].normalized_text


@pytest.mark.parametrize("missing", ["relationship", "member", "external"])
def test_unresolved_referenced_word_header_is_not_complete(missing):
    import zipfile

    from lxml import etree

    word = WordDocument()
    word.add_paragraph("Readable body survives")
    word.sections[0].header.paragraphs[0].text = "Referenced header"
    output = io.BytesIO()
    word.save(output)
    changed = io.BytesIO()
    with zipfile.ZipFile(output) as source, zipfile.ZipFile(changed, "w") as target:
        rels = etree.fromstring(source.read("word/_rels/document.xml.rels"))
        header = next(row for row in rels if row.get("Type", "").endswith("/header"))
        member = "word/" + header.get("Target")
        if missing == "relationship":
            rels.remove(header)
        elif missing == "external":
            header.set("TargetMode", "External")
            header.set("Target", "https://not-loaded.invalid/header.xml")
        for name in source.namelist():
            if missing == "member" and name == member:
                continue
            target.writestr(name, etree.tostring(rels) if name == "word/_rels/document.xml.rels" else source.read(name))
    batch = reader()(changed, "missing-header.docx", policy=ExtractionPolicy())
    assert "Readable body survives" in "\n".join(unit.normalized_text for unit in batch.units)
    assert batch.coverage != "complete"
    assert batch.omitted_reasons or any(unit.state == "action_required" for unit in batch.units)


def test_spreadsheet_cached_formula_gap_and_coordinates():
    workbook = Workbook()
    workbook.active["A1"] = "Context"
    workbook.active["B2"] = "=SUM(1,2)"
    output = io.BytesIO()
    workbook.save(output)
    batch = reader()(output, "formula.xlsx", policy=ExtractionPolicy())
    assert "A1" in batch.units[0].normalized_text
    assert "Context" in batch.units[0].normalized_text
    assert batch.coverage != "complete"
    assert any("cache" in note.lower() for unit in batch.units for note in unit.warnings)


def test_workbook_embedded_code_is_never_claimed_as_read():
    import zipfile

    workbook = Workbook()
    workbook.active["A1"] = "Readable cell"
    output = io.BytesIO()
    workbook.save(output)
    with zipfile.ZipFile(output, "a") as archive:
        archive.writestr("xl/vbaProject.bin", b"unexecuted fixture macro bytes")
    batch = reader()(output, "macro.xlsm", policy=ExtractionPolicy())
    assert batch.coverage != "complete"
    assert any("macro" in note.lower() for unit in batch.units for note in unit.warnings)


def test_text_splits_units_without_discarding_late_content():
    policy = replace(ExtractionPolicy(), max_unit_chars=40)
    batch = reader()(io.BytesIO(b"beginning " * 20 + b"late appendix phrase"), "long.txt", policy=policy)
    assert len(batch.units) > 1
    assert "late appendix phrase" in "".join(unit.raw_text for unit in batch.units)
    assert all(len(unit.raw_text) <= 40 for unit in batch.units)
    assert batch.coverage == "complete"


@pytest.mark.parametrize("extension", ["doc", "xls", "ppt", "pptx"])
def test_unsupported_accepted_format_is_explicit(extension):
    batch = reader()(io.BytesIO(b"unsupported format"), f"old.{extension}", policy=ExtractionPolicy())
    assert batch.units[0].state == "unsupported"
    assert batch.coverage != "complete"


def test_compressed_image_pixel_limit_is_bounded():
    output = io.BytesIO()
    Image.new("RGB", (500, 500), "white").save(output, "PNG")
    batch = reader()(output, "scan.png", policy=replace(ExtractionPolicy(), max_image_pixels=1000))
    assert batch.coverage != "complete"
    assert any("pixel" in note.lower() for unit in batch.units for note in unit.warnings)


def test_frame_limit_does_not_require_full_file_frame_count():
    from PIL.TiffImagePlugin import TiffImageFile

    output = io.BytesIO()
    frames = [Image.new("RGB", (30, 30), color) for color in ("white", "gray", "black")]
    frames[0].save(output, "TIFF", save_all=True, append_images=frames[1:])

    def unavailable_count(_image):
        raise ValueError("Counting every frame is not a bounded operation")

    with patch.object(TiffImageFile, "n_frames", property(unavailable_count)):
        batch = reader()(output, "frames.tiff", policy=replace(ExtractionPolicy(), max_frames=2))
    assert [unit.locator_key for unit in batch.units] == ["image:frame:1", "image:frame:2"]
    assert batch.expected_units is None
    assert batch.omitted_reasons
