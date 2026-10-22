"""Bounded native source readers. No database mutations or provider requests."""
from __future__ import annotations

import io
import math
import posixpath
import zipfile

from apps.core.utils import file_extension, normalise_text

from .processing_contracts import Coverage, ExtractionPolicy, UnitBatch, UnitResult, WorkState

IMAGE_TYPES = {"jpg", "jpeg", "png", "tif", "tiff", "bmp", "gif", "webp"}
TEXT_TYPES = {"txt", "csv", "md", "log"}


def _unit(key, ordinal, kind, text="", *, policy, warnings=(), metadata=None, state=WorkState.COMPLETE):
    truncated = len(text) > policy.max_unit_chars
    if truncated:
        warnings = (*warnings, "Source-unit text reached the storage safety limit; later units are retained.")
    raw = text[:policy.max_unit_chars]
    return UnitResult(locator_key=key, ordinal=ordinal, kind=kind, raw_text=raw, normalized_text=normalise_text(raw),
        state=state, coverage=Coverage.COMPLETE if state == WorkState.COMPLETE and not warnings else Coverage.PARTIAL,
        warnings=tuple(warnings), truncated=truncated, metadata=metadata or {})


def _failure(note, *, policy, state=WorkState.FAILED, locator="file"):
    unit = _unit(locator, 1, "file", policy=policy, state=state, warnings=(note,))
    return UnitBatch(units=(unit,), expected_units=None, omitted_reasons=(note,), policy_digest=policy.digest)


def _pdf(payload, policy):
    from pypdf import PdfReader, apply_configuration

    from .pdf_assessment import PDFIUM_LOCK, assess_pdf_page

    units = []
    limit = min(policy.max_archive_bytes, 32 * 1024 * 1024)
    with apply_configuration(zlib_maximum_output_length=limit, lzw_maximum_output_length=limit,
        run_length_maximum_output_length=limit, maximum_declared_stream_length=limit,
        array_based_stream_maximum_output_length=limit, image_maximum_buffer_size=limit,
        jbig2dec_binary=None, xform_maximum_invocations_per_extraction=1000):
        reader = PdfReader(io.BytesIO(payload), strict=False)
        if reader.is_encrypted:
            return _failure("Encrypted PDF content cannot be read automatically.", policy=policy, state=WorkState.ACTION_REQUIRED)
        expected = len(reader.pages)
        for index, page in enumerate(reader.pages[:policy.max_pages]):
            warnings = []
            raster = None
            try:
                width, height = float(page.mediabox.width), float(page.mediabox.height)
                pixels = math.ceil(width * policy.render_dpi/72) * math.ceil(height * policy.render_dpi/72)
                if pixels <= 0 or pixels > policy.max_image_pixels:
                    warnings.append("Rendering pixel limit reached; the page was not silently downsampled.")
                else:
                    try:
                        import pypdfium2
                        with PDFIUM_LOCK, pypdfium2.PdfDocument(payload) as pdf:
                            rendered = pdf[index]
                            try:
                                bitmap = rendered.render(scale=policy.render_dpi/72)
                                try:
                                    raster = bitmap.to_pil().copy()
                                finally:
                                    bitmap.close()
                            finally:
                                rendered.close()
                    except (ImportError, RuntimeError, OSError, ValueError, TypeError):
                        warnings.append("The PDF renderer is unavailable; native text was preserved.")
                assessment = assess_pdf_page(page, raster, policy=policy)
                warnings.extend(assessment.pop("warnings"))
                text = assessment.pop("text")
                if assessment["ocr_needed"]:
                    warnings.append("This page needs OCR or review of unread visual content.")
                unit = _unit(f"pdf:page:{index+1}", index+1, "pdf_page", text,
                    policy=policy, warnings=warnings, metadata={**assessment, "page_number": index+1})
                units.append(unit)
            except Exception:
                units.append(_unit(f"pdf:page:{index+1}", index+1, "pdf_page", policy=policy,
                    state=WorkState.ACTION_REQUIRED, warnings=("This PDF page could not be read within safe limits.",)))
            finally:
                if raster is not None:
                    raster.close()
    omitted = ("PDF page processing limit reached; additional pages remain unread.",) if expected > len(units) else ()
    return UnitBatch(tuple(units), expected, omitted, policy.digest)


def _archive(payload, policy):
    archive = zipfile.ZipFile(io.BytesIO(payload))
    members = archive.infolist()
    if (len(members) > policy.max_archive_members or len({member.filename for member in members}) != len(members)
        or sum(member.file_size for member in members) > policy.max_archive_bytes
        or any(member.flag_bits & 1 for member in members)):
        archive.close()
        raise ValueError("Office archive exceeds safe limits")
    return archive


def _xml(archive, name, policy):
    from lxml import etree

    if archive.getinfo(name).file_size > min(policy.max_archive_bytes, 16*1024*1024):
        raise ValueError("Office XML part exceeds safe limits")
    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    return etree.fromstring(archive.read(name), parser=parser)


def _word_limit(units, policy):
    return UnitBatch(tuple(units), None, ("Word source-unit safety limit reached; remaining content is unmeasured.",), policy.digest)


def _word(payload, policy):
    units, omitted = [], []
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    with _archive(payload, policy) as archive:
        names = set(archive.namelist())

        def relationships(part):
            location = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part)+".rels")
            if location not in names:
                return {}
            return {row.get("Id"): row for row in _xml(archive, location, policy)}

        def target(part, rel):
            if rel is None or rel.get("TargetMode") == "External":
                return None
            path = rel.get("Target", "")
            resolved = posixpath.normpath(path.lstrip("/") if path.startswith("/") else posixpath.join(posixpath.dirname(part), path))
            return resolved if resolved in names and not resolved.startswith("../") else None

        main = _xml(archive, "word/document.xml", policy)
        main_rels = relationships("word/document.xml")
        parts = [("word/document.xml", main)]
        for note_part in ("word/footnotes.xml", "word/endnotes.xml", "word/comments.xml"):
            if note_part in names:
                parts.append((note_part, _xml(archive, note_part, policy)))
        for ref in main.xpath("//w:headerReference | //w:footerReference", namespaces=ns):
            path = target("word/document.xml", main_rels.get(ref.get(f"{{{ns['r']}}}id")))
            if not path:
                if len(units) >= policy.max_word_units:
                    return _word_limit(units, policy)
                note = "A referenced Word header/footer is unavailable; external resources were not loaded."
                omitted.append(note)
                units.append(_unit(f"docx:unread-header-footer:{len(units)+1}", len(units)+1,
                    "unread_header_footer", policy=policy, state=WorkState.ACTION_REQUIRED, warnings=(note,)))
            if path and path not in {entry[0] for entry in parts}:
                parts.append((path, _xml(archive, path, policy)))
        for part, tree in parts:
            rels = relationships(part)
            blocks = tree.find("w:body", ns) if part == "word/document.xml" else tree
            if blocks is None:
                continue
            for block_index, block in enumerate(blocks, 1):
                kind = block.tag.rsplit("}", 1)[-1]
                if kind == "sectPr":
                    continue
                if len(units) >= policy.max_word_units:
                    return _word_limit(units, policy)
                text = "\n".join(" ".join(p.itertext(f"{{{ns['w']}}}t")) for p in block.iter(f"{{{ns['w']}}}p"))
                warnings = []
                if any(node.tag.rsplit("}", 1)[-1] in {"altChunk", "object", "pict", "txbxContent"} for node in block.iter()):
                    warnings.append("Unsupported Word object or alternate content may remain unread.")
                if any(node.tag.rsplit("}", 1)[-1] in {"chart", "relIds", "contentPart"} for node in block.iter()):
                    warnings.append("Word chart, diagram or linked content remains incompletely read.")
                if kind not in {"p", "tbl"}:
                    warnings.append("Unsupported Word block may remain unread.")
                key = f"docx:{part}:block:{block_index}"
                units.append(_unit(key, len(units)+1, "table" if kind == "tbl" else "paragraph", text,
                    policy=policy, warnings=warnings, metadata={"archive_part": part, "block": block_index}))
                for image_index, drawing in enumerate(block.xpath(".//*[local-name()='blip']"), 1):
                    if len(units) >= policy.max_word_units:
                        return _word_limit(units, policy)
                    rel = rels.get(drawing.get(f"{{{ns['r']}}}embed"))
                    path = target(part, rel)
                    note = "Embedded image needs OCR; it has not been read as text."
                    if not path:
                        note = "An embedded or linked image is unavailable; remote resources were not loaded."
                    units.append(_unit(f"{key}:image:{image_index}", len(units)+1, "embedded_image", policy=policy,
                        state=WorkState.PENDING if path else WorkState.ACTION_REQUIRED,
                        warnings=(note,), metadata={"archive_member": path, "ocr_needed": bool(path)}))
        for member in sorted(name for name in names if name.startswith(("word/embeddings/", "word/diagrams/", "word/charts/"))):
            if len(units) >= policy.max_word_units:
                return _word_limit(units, policy)
            units.append(_unit(f"docx:unsupported:{len(units)+1}", len(units)+1, "embedded_object",
                policy=policy, state=WorkState.UNSUPPORTED, warnings=("Embedded Word object, chart or diagram remains unread.",),
                metadata={"archive_member": member}))
    return UnitBatch(tuple(units), None if omitted else len(units), tuple(dict.fromkeys(omitted)), policy.digest)


def _xlsx(payload, policy):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    with _archive(payload, policy) as archive:
        embedded = [member for member in archive.namelist() if member.startswith(("xl/media/", "xl/embeddings/", "xl/charts/", "xl/externalLinks/")) or member.lower().endswith("vbaproject.bin")]
    formulas = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    units, omitted = [], []
    cells = 0
    try:
        for sheet, values in zip(formulas.worksheets, cached.worksheets, strict=True):
            rows, columns = min(sheet.max_row or 1, policy.max_rows), min(sheet.max_column or 1, policy.max_columns)
            if (sheet.max_row or 1) > rows or (sheet.max_column or 1) > columns:
                omitted.append("Worksheet row/column safety limit reached; additional cells remain unread.")
            lines, warnings, start = [], [], 1
            for row_index, (row, cache_row) in enumerate(zip(sheet.iter_rows(max_row=rows, max_col=columns),
                values.iter_rows(max_row=rows, max_col=columns), strict=True), 1):
                if cells + columns > policy.max_cells:
                    omitted.append("Total spreadsheet cell safety limit reached.")
                    break
                cells += columns
                line = []
                for column, (cell, cached_cell) in enumerate(zip(row, cache_row, strict=True), 1):
                    value = cached_cell.value if cell.data_type == "f" else cell.value
                    coordinate = f"{get_column_letter(column)}{row_index}"
                    if cell.data_type == "f" and value is None:
                        warnings.append(f"{coordinate}: formula cache unavailable; formulas are never executed.")
                        value = cell.value
                    if value not in (None, ""):
                        line.append(f"{coordinate}: {value}")
                if line:
                    lines.append(" | ".join(line))
                if row_index % 100 == 0 or row_index == rows:
                    units.append(_unit(f"xlsx:{sheet.title}:rows:{start}-{row_index}", len(units)+1, "sheet_range",
                        f"{sheet.title}\n" + "\n".join(lines), policy=policy, warnings=warnings,
                        metadata={"sheet": sheet.title, "start_row": start, "end_row": row_index}))
                    lines, warnings, start = [], [], row_index+1
            if lines or warnings:
                units.append(_unit(f"xlsx:{sheet.title}:rows:{start}-{row_index}", len(units)+1, "sheet_range",
                    f"{sheet.title}\n" + "\n".join(lines), policy=policy, warnings=warnings,
                    metadata={"sheet": sheet.title, "start_row": start, "end_row": row_index}))
            if cells >= policy.max_cells:
                if sheet != formulas.worksheets[-1]:
                    omitted.append("Additional worksheets remain unread because of the cell limit.")
                break
        for index, member in enumerate(embedded, 1):
            is_image = member.startswith("xl/media/")
            units.append(_unit(f"xlsx:embedded:{index}", len(units)+1, "embedded_image" if is_image else "embedded_object",
                policy=policy, state=WorkState.PENDING if is_image else WorkState.UNSUPPORTED,
                warnings=("Spreadsheet macros are never executed or read." if member.lower().endswith("vbaproject.bin") else "Embedded spreadsheet image, object or linked content remains unread.",), metadata={"archive_member": member, "ocr_needed": is_image}))
    finally:
        formulas.close()
        cached.close()
    return UnitBatch(tuple(units), None if omitted else len(units), tuple(dict.fromkeys(omitted)), policy.digest)


def _images(payload, policy):
    from PIL import Image

    units, omitted = [], []
    with Image.open(io.BytesIO(payload)) as image:
        expected = None
        for index in range(policy.max_frames):
            try:
                image.seek(index)
            except EOFError:
                expected = index
                break
            too_large = image.width * image.height > policy.max_image_pixels
            note = "Image pixel limit reached; no silent downsampling was performed." if too_large else "This frame needs OCR; near-blank appearance does not verify empty content."
            units.append(_unit(f"image:frame:{index+1}", index+1, "image_frame", policy=policy,
                state=WorkState.ACTION_REQUIRED if too_large else WorkState.PENDING, warnings=(note,),
                metadata={"frame": index+1, "width": image.width, "height": image.height, "ocr_needed": not too_large}))
        if expected is None:
            try:
                image.seek(policy.max_frames)
            except EOFError:
                expected = len(units)
            else:
                omitted.append("Image frame limit reached; additional frame coverage is unknown.")
    return UnitBatch(tuple(units), expected, tuple(omitted), policy.digest)


def read_source_units(file_obj, filename: str, *, policy: ExtractionPolicy) -> UnitBatch:
    """Original bytes remain untouched. Parser failures are safe, visible outcomes."""
    try:
        file_obj.seek(0)
        payload = file_obj.read(policy.max_input_bytes+1)
        if len(payload) > policy.max_input_bytes:
            return _failure("File input safety limit reached.", policy=policy, state=WorkState.ACTION_REQUIRED)
        extension = file_extension(filename)
        if extension == "pdf":
            return _pdf(payload, policy)
        if extension in {"docx", "dotx"}:
            return _word(payload, policy)
        if extension in {"xlsx", "xlsm"}:
            return _xlsx(payload, policy)
        if extension in IMAGE_TYPES:
            return _images(payload, policy)
        if extension in TEXT_TYPES:
            text = payload.decode("utf-8", errors="replace")
            warning = ("Some bytes were not valid UTF-8; text quality needs review.",) if "\ufffd" in text else ()
            units = tuple(_unit(f"text:offset:{offset}", index+1, "text_range", text[offset:offset+policy.max_unit_chars],
                policy=policy, warnings=warning, metadata={"offset": offset})
                for index, offset in enumerate(range(0, max(1, len(text)), policy.max_unit_chars)))
            return UnitBatch(units, len(units), (), policy.digest)
        return _failure(f"Automatic text extraction is unsupported for .{extension}; use manual metadata and the original file.",
            policy=policy, state=WorkState.UNSUPPORTED)
    except ImportError:
        return _failure("A required native reader dependency is unavailable.", policy=policy, state=WorkState.ACTION_REQUIRED)
    except Exception:
        return _failure("This file could not be read within safe parsing limits. The original file is unchanged.", policy=policy)
    finally:
        file_obj.seek(0)
