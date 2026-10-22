"""Bounded, read-only previews. Derived DOCX bytes are never stored."""

import csv
import io
import posixpath
import struct
import zipfile
from xml.etree import ElementTree as ET

from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile

from apps.core.utils import file_extension, validate_upload

MIME_TYPES = {
    "pdf": "application/pdf", "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "txt": "text/plain; charset=utf-8",
    "csv": "text/csv; charset=utf-8", "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
TEXT_LIMIT = 200_000
ROW_LIMIT = 200
COL_LIMIT = 50
CELL_LIMIT = 2000
IMAGE_PIXEL_LIMIT = 16_000_000
# openpyxl may use lxml, whose XMLSyntaxError is a SyntaxError, not ET.ParseError.
PREVIEW_ERRORS = (OSError, ValueError, ValidationError, KeyError, IndexError, TypeError, SyntaxError, zipfile.BadZipFile, csv.Error)


def read_validated(item):
    limit = settings.MAX_UPLOAD_MB * 1024 * 1024
    with item.file.open("rb") as source:
        data = source.read(limit + 1)
    if len(data) > limit:
        raise ValidationError("This file exceeds the preview size limit. Download the original to view it.")
    validate_upload(ContentFile(data, name=item.original_name))
    return data


def safe_xml(data):
    # Reject entity declarations before parsing. OpenXML has no need for a DTD.
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValidationError("This document contains unsupported XML declarations.")
    return ET.fromstring(data)


def image_pixels(data, extension):
    """Read bounded raster headers before the browser allocates decoded pixels."""
    signatures = {
        "png": data.startswith(b"\x89PNG\r\n\x1a\n"),
        "jpg": data.startswith(b"\xff\xd8\xff"), "jpeg": data.startswith(b"\xff\xd8\xff"),
        "gif": data.startswith((b"GIF87a", b"GIF89a")), "bmp": data.startswith(b"BM"),
    }
    if not signatures.get(extension):
        raise ValidationError("Embedded image content does not match its filename. Download the original.")
    width = height = 0
    if extension == "png" and len(data) >= 24 and data[12:16] == b"IHDR":
        width, height = struct.unpack(">II", data[16:24])
    elif extension in {"jpg", "jpeg"}:
        position = 2
        while position + 4 < min(len(data), 1_000_000):
            if data[position] != 0xFF:
                break
            while position < len(data) and data[position] == 0xFF:
                position += 1
            if position >= len(data):
                break
            marker = data[position]
            position += 1
            if marker in {0xD8, 0xD9, 0x01, *range(0xD0, 0xD8)}:
                continue
            if position + 2 > len(data) or marker == 0xDA:
                break
            length = int.from_bytes(data[position:position + 2], "big")
            if length < 2 or position + length > len(data):
                break
            if marker in {0xC0, 0xC1, 0xC2, 0xC3, 0xC5, 0xC6, 0xC7, 0xC9, 0xCA, 0xCB, 0xCD, 0xCE, 0xCF} and length >= 8:
                height, width = struct.unpack(">HH", data[position + 3:position + 7])
                break
            position += length
    elif extension == "gif" and len(data) >= 10:
        width, height = struct.unpack("<HH", data[6:10])
    elif extension == "bmp" and len(data) >= 26:
        header_size = int.from_bytes(data[14:18], "little")
        if header_size >= 40:
            width, height = struct.unpack("<ii", data[18:26])
            height = abs(height)
        elif header_size == 12:
            width, height = struct.unpack("<HH", data[18:22])
    if width <= 0 or height <= 0 or width > 16384 or height > 16384 or width * height > IMAGE_PIXEL_LIMIT:
        raise ValidationError("Image dimensions exceed the safe preview limit or cannot be read. Download the original.")
    return width * height


def sanitized_docx(data):
    """Drop every relationship except local XML, raster images and styles.

    Keeping only raster embedded images also prevents external URLs in SVG or
    HTML parts from reaching the renderer. Unknown/active parts are omitted.
    """
    output = io.BytesIO()
    rendered_pixels = 0
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        allowed = set()
        for name in source.namelist():
            lower = name.lower()
            if lower.startswith("/") or ".." in lower.split("/") or "\\" in lower:
                raise ValidationError("This document has an invalid package path.")
            if lower.endswith((".xml", ".rels")) or lower.endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp")):
                if not any(part in lower for part in ("activex", "embeddings/", "vba", "customxml/")):
                    allowed.add(name)
        for name in allowed:
            payload = source.read(name)
            extension = file_extension(name)
            if extension in {"png", "jpg", "jpeg", "gif", "bmp"}:
                rendered_pixels += image_pixels(payload, extension)
                if rendered_pixels > IMAGE_PIXEL_LIMIT * 2:
                    raise ValidationError("Embedded images exceed the document preview limit. Download the original.")
            if name.endswith((".xml", ".rels")):
                root = safe_xml(payload)
                if name.endswith(".rels"):
                    base = posixpath.dirname(posixpath.dirname(name))
                    for relationship in list(root):
                        destination = relationship.get("Target", "")
                        relation = relationship.get("Type", "").rsplit("/", 1)[-1].lower()
                        resolved = posixpath.normpath(posixpath.join(base, destination))
                        if (
                            relationship.get("TargetMode", "").lower() == "external"
                            or ":" in destination or destination.startswith(("/", "\\"))
                            or resolved not in allowed
                            or relation not in {"officedocument", "styles", "numbering", "settings", "theme", "image", "footnotes", "endnotes", "header", "footer", "fonttable"}
                        ):
                            root.remove(relationship)
                else:
                    for element in root.iter():
                        for child in list(element):
                            if child.tag.rsplit("}", 1)[-1] in {"altChunk", "object", "OLEObject", "control"}:
                                element.remove(child)
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(name, payload)
    return output.getvalue()


def tabular_preview(data, extension, worksheet=None):
    truncated = False
    sheets = []
    if extension == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheets = [sheet.title for sheet in workbook.worksheets]
            if not sheets:
                raise ValidationError("This workbook has no worksheets to preview.")
            if worksheet is not None and worksheet not in sheets:
                raise ValidationError("Choose a worksheet from this file.")
            worksheet = worksheet or sheets[0]
            sheet = workbook[worksheet]
            # Ignore attacker-controlled worksheet dimensions; read a bounded iterator.
            sheet.reset_dimensions()
            values = sheet.iter_rows(max_row=ROW_LIMIT + 1, max_col=COL_LIMIT + 1, values_only=True)
            rows = []
            for index, row in enumerate(values):
                if index == ROW_LIMIT:
                    truncated = True
                    break
                truncated |= any(value is not None for value in row[COL_LIMIT:])
                truncated |= any(len(str(value)) > CELL_LIMIT for value in row[:COL_LIMIT] if value is not None)
                rows.append([str(value)[:CELL_LIMIT] if value is not None else "" for value in row[:COL_LIMIT]])
        finally:
            workbook.close()
    else:
        text = data[:TEXT_LIMIT].decode("utf-8-sig", errors="replace")
        truncated = len(data) > TEXT_LIMIT
        rows = []
        # The entire input is bounded, including a maliciously long single field.
        for index, row in enumerate(csv.reader(io.StringIO(text))):
            if index == ROW_LIMIT:
                truncated = True
                break
            truncated |= len(row) > COL_LIMIT or any(len(value) > CELL_LIMIT for value in row)
            rows.append([value[:CELL_LIMIT] for value in row[:COL_LIMIT]])
    return {"rows": rows, "truncated": truncated, "worksheets": sheets, "worksheet": worksheet}


def preview_kind(filename):
    extension = file_extension(filename)
    return {"jpg": "image", "jpeg": "image", "png": "image", "csv": "table", "xlsx": "table"}.get(
        extension, extension if extension in {"pdf", "docx", "txt"} else "unavailable"
    )
